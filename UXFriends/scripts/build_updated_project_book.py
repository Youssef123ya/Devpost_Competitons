from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_PATH = Path(
    "/Users/hendawiy/Documents/project/UXFriends/docs/UXFriends-Updated-Project-Documentation.docx"
)

ACCENT = RGBColor(110, 46, 196)
TEAL = RGBColor(24, 165, 149)
NAVY = RGBColor(18, 24, 45)
GRAY = RGBColor(89, 95, 112)
LIGHT_GRAY = RGBColor(245, 246, 250)
WHITE = RGBColor(255, 255, 255)


def set_cell_shading(cell, fill: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell_properties.append(shading)


def set_table_borders(table, color: str = "D8DDE8") -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "4")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def set_page_margins(section) -> None:
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)


def setup_document() -> Document:
    document = Document()
    set_page_margins(document.sections[0])

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = NAVY

    heading1 = styles["Heading 1"]
    heading1.font.name = "Calibri"
    heading1.font.size = Pt(19)
    heading1.font.bold = True
    heading1.font.color.rgb = NAVY

    heading2 = styles["Heading 2"]
    heading2.font.name = "Calibri"
    heading2.font.size = Pt(14)
    heading2.font.bold = True
    heading2.font.color.rgb = ACCENT

    heading3 = styles["Heading 3"]
    heading3.font.name = "Calibri"
    heading3.font.size = Pt(12.5)
    heading3.font.bold = True
    heading3.font.color.rgb = TEAL

    return document


def add_cover_page(document: Document) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run("UXFriends").font.color.rgb = ACCENT

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_before = Pt(68)
    p.space_after = Pt(8)
    run = p.add_run("UXFriends")
    run.font.name = "Calibri"
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = NAVY

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Updated Project Documentation")
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = ACCENT

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_after = Pt(16)
    run = p.add_run(
        "Implementation-aligned revision of the original graduation book"
    )
    run.font.size = Pt(12)
    run.font.italic = True
    run.font.color.rgb = GRAY

    callout = document.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    callout.autofit = True
    set_table_borders(callout, "D9C8FF")
    cell = callout.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "F8F4FF")
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.space_before = Pt(10)
    para.space_after = Pt(10)
    run = para.add_run(
        "This revision reflects the real UXFriends application as implemented on April 28, 2026. "
        "It replaces outdated assumptions in the old PDF such as SQL Server as the current default, "
        "mandatory cloud deployment, and AI-generated quizzes as a finished feature."
    )
    run.font.size = Pt(11.5)
    run.font.color.rgb = NAVY

    document.add_paragraph("")
    info = document.add_table(rows=4, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.columns[0].width = Inches(2.0)
    info.columns[1].width = Inches(4.6)
    set_table_borders(info)
    rows = [
        ("Application", "UXFriends"),
        ("Current stack", "React/Vite frontend, ASP.NET Core backend, SQLite by default, FastAPI AI service"),
        ("Document purpose", "Update the original project book to match the current codebase, features, removals, and business rules"),
        ("Revision date", date.today().strftime("%B %d, %Y")),
    ]
    for idx, (label, value) in enumerate(rows):
        left = info.cell(idx, 0)
        right = info.cell(idx, 1)
        set_cell_shading(left, "F0E9FF")
        for paragraph in (left.paragraphs[0], right.paragraphs[0]):
            paragraph.space_before = Pt(5)
            paragraph.space_after = Pt(5)
        left.paragraphs[0].add_run(label).bold = True
        left.paragraphs[0].runs[0].font.color.rgb = NAVY
        right.paragraphs[0].add_run(value).font.color.rgb = NAVY

    document.add_paragraph("")
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Prepared as an implementation-aligned documentation update for graduation submission."
    )
    run.font.size = Pt(11)
    run.font.color.rgb = GRAY

    document.add_page_break()


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def add_paragraph(document: Document, text: str, *, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        first.font.color.rgb = NAVY
        rest = paragraph.add_run(text[len(bold_prefix) :])
        rest.font.color.rgb = NAVY
    else:
        run = paragraph.add_run(text)
        run.font.color.rgb = NAVY


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(item)
        run.font.color.rgb = NAVY


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(item)
        run.font.color.rgb = NAVY


def add_table(document: Document, headers: list[str], rows: list[list[str]], *, first_col_fill="F7F3FF") -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    set_table_borders(table)
    for idx, title in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "EEE7FF")
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.space_before = Pt(3)
        para.space_after = Pt(3)
        run = para.add_run(title)
        run.bold = True
        run.font.color.rgb = NAVY

    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cell = cells[idx]
            if idx == 0 and first_col_fill:
                set_cell_shading(cell, first_col_fill)
            para = cell.paragraphs[0]
            para.space_before = Pt(3)
            para.space_after = Pt(3)
            run = para.add_run(value)
            run.font.size = Pt(10.5)
            run.font.color.rgb = NAVY

    document.add_paragraph("")


def add_callout(document: Document, title: str, body: str, fill: str = "F5FAFF") -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, "C9DAF8")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.space_before = Pt(8)
    p.space_after = Pt(8)
    title_run = p.add_run(f"{title}: ")
    title_run.bold = True
    title_run.font.color.rgb = NAVY
    body_run = p.add_run(body)
    body_run.font.color.rgb = NAVY
    document.add_paragraph("")


def add_manual_toc(document: Document) -> None:
    add_heading(document, "Table of Contents", level=1)
    sections = [
        "Executive Abstract",
        "Chapter 1: Introduction and Update Rationale",
        "Chapter 2: Current Implemented System",
        "Chapter 3: Data, APIs, and Business Rules",
        "Chapter 4: Feature Status Matrix",
        "Chapter 5: Quality, Operations, and Known Limitations",
        "Chapter 6: Diagram and Graduation Guidance",
        "Appendix A: Current Backend Modules",
        "Appendix B: Current Frontend Pages and Route Groups",
        "Appendix C: Added, Removed, and Deferred Scope",
    ]
    add_numbered(document, sections)
    document.add_page_break()


def add_content(document: Document) -> None:
    add_heading(document, "Executive Abstract", level=1)
    add_paragraph(
        document,
        "UXFriends is now a fully implemented web application for learning UX/UI design through structured pathways, quizzes, projects, gamification, and an AI-supported coaching layer. The current system runs as a React/Vite frontend, an ASP.NET Core 8 backend API, a FastAPI AI service, and a relational data layer managed through EF Core. In the local development environment, SQLite is the active default database provider, while SQL Server remains an optional provider rather than the current source of truth."
    )
    add_paragraph(
        document,
        "Compared with the original project book, the current application now includes optional email verification with a 48-hour deadline, SMTP-based password reset, cookie-based authentication, a one-time onboarding placement assessment, persisted lesson notes, section certificates, public/private portfolio pages, daily goals, level-based challenges, weekly leagues, and persisted AI chat sessions. At the same time, some ideas from the old document are no longer accurate as implemented facts. The current application does not include a human reviewer workflow for projects, does not expose a student shop or coins economy, and does not currently rely on SQL Server as its active default database in local use."
    )
    add_callout(
        document,
        "Scope note",
        "This updated document is implementation-aligned. When the old report described planned cloud deployment, AI-generated quizzes, or SQL Server as the live database, those statements are revised here to reflect the real codebase as currently shipped.",
    )

    add_heading(document, "Chapter 1: Introduction and Update Rationale", level=1)
    add_heading(document, "1.1 Problem domain", level=2)
    add_paragraph(
        document,
        "UXFriends belongs to the educational technology domain and specifically targets beginner-to-advanced learning in UX/UI design. The product addresses a common gap in self-learning: scattered resources, low accountability, weak progression systems, and limited practical application. Rather than leaving the learner to assemble theory, practice, and feedback from separate sources, UXFriends brings those layers into a single web platform."
    )
    add_heading(document, "1.2 Why the original book needed revision", level=2)
    add_bullets(
        document,
        [
            "The old book treated some design goals as if they were already implemented, while the real codebase now contains a more concrete but also more specific feature set.",
            "The old text referred to Microsoft SQL Server as the system database, but the current application runs on SQLite by default in local development with EF Core migration support.",
            "The old text described AI-generated quizzes and periodic report generation as core implemented AI features, while the real AI layer currently focuses on mistake-based analysis, revision planning, recommendations, and persisted coaching chat.",
            "The product scope changed during implementation: the shop and coin system were removed, portfolio privacy was simplified to public/private, and projects are auto-approved because no reviewer role exists in the shipped application.",
        ],
    )
    add_heading(document, "1.3 Current product snapshot", level=2)
    add_table(
        document,
        ["Area", "Current implementation"],
        [
            ["Curriculum", "3 levels (Beginner, Intermediate, Advanced) with 15 sections, 30 units, 90 lessons"],
            ["Assessment", "Lesson quizzes, unit quizzes, section final exams, one-time placement test"],
            ["Projects", "15 capstone projects, auto-approved on valid submission"],
            ["Gamification", "XP, streaks, badges, daily goals, level-based challenges, weekly leagues"],
            ["AI support", "FastAPI AI coach with heuristic/local provider modes, revision plan, mistake analysis, persisted chat sessions"],
            ["Profile and portfolio", "Editable profile, badges, progress, public/private portfolio, public portfolio viewing from leaderboard"],
            ["Account flows", "Register, login, logout, forgot/reset password, optional email verification, Google/GitHub OAuth if configured"],
        ],
    )

    add_heading(document, "Chapter 2: Current Implemented System", level=1)
    add_heading(document, "2.1 Technology stack", level=2)
    add_table(
        document,
        ["Layer", "Current stack", "Implementation notes"],
        [
            ["Frontend", "React 19 + Vite + TypeScript", "Single-page application with route-based page switching and Arabic/English language support"],
            ["Backend API", "ASP.NET Core 8", "REST controllers for auth, learning, gamification, portfolio, search, AI, and public data"],
            ["Database", "SQLite by default, SQL Server optional", "EF Core is the active data access layer; local development now runs migration-based startup"],
            ["AI service", "FastAPI", "Supports heuristic mode and local-model mode, with persisted chat history in the main backend database"],
            ["Authentication", "ASP.NET Identity + HttpOnly auth cookie", "JWT is generated server-side and stored in a secure cookie rather than browser local storage"],
            ["Email", "SMTP", "Used for password reset and optional account email verification"],
        ],
    )
    add_heading(document, "2.2 Runtime service architecture", level=2)
    add_bullets(
        document,
        [
            "The React frontend at `http://localhost:5173` consumes the backend through an Axios API client configured with `withCredentials` for cookie-based authentication.",
            "The ASP.NET Core backend at `http://localhost:5176` owns users, progress, gamification, portfolio, certificates, and all persistence-heavy business rules.",
            "The FastAPI service at `http://localhost:8001` provides AI coaching behavior. It can run in heuristic mode or in local-model mode through an Ollama-compatible endpoint.",
            "The backend database stores identity, progress, quiz attempts, XP logs, challenges, badges, email verification tokens, password reset tokens, and AI chat session history.",
        ],
    )
    add_callout(
        document,
        "Current source of truth",
        "The backend API is the operational source of truth for profile state, progress, XP, streaks, badges, notes, certificates, challenge status, and leaderboard data. The frontend renders these values and only keeps lightweight session cache for responsiveness.",
    )
    add_heading(document, "2.3 Current functional modules", level=2)
    add_table(
        document,
        ["Module", "Implemented behavior"],
        [
            ["Authentication and account lifecycle", "Register, login, logout, OAuth callback flow, forgot/reset password, optional email verification, protected routes, cookie auth"],
            ["Onboarding and level determination", "Welcome page, start-from-scratch path, one-time 18-question placement assessment, level persistence, level-based content access"],
            ["Learning content", "Sections, units, lessons, lesson notes, lesson completion tracking, lesson quizzes, unit quizzes, section finals"],
            ["Projects", "Section-level project briefs and submissions, auto-approval, XP reward, publication into portfolio"],
            ["Gamification", "XP, streaks, badges, daily goals, challenge join/claim flow, weekly leagues and rankings"],
            ["AI and revision", "AI coach overview, analysis page, recommendation cards, revision plans, revision sessions, rolling chat persistence"],
            ["Profile and portfolio", "Profile statistics, badges tab, progress overview, public/private portfolio page, public access from leaderboard"],
            ["Search and discovery", "Search endpoint with typed results for sections, lessons, projects, and certificates"],
        ],
    )
    add_heading(document, "2.4 Authentication and account rules", level=2)
    add_bullets(
        document,
        [
            "Registration creates a user account immediately; email verification is optional but strongly encouraged and visible inside the app.",
            "If an account remains unverified for 48 hours after creation, it can be deleted automatically together with its progress data.",
            "Authentication is no longer stored in browser local storage as a JWT token. The application uses an `HttpOnly` cookie for safer session handling.",
            "The remember-me option changes cookie persistence rather than storing a bearer token in client storage.",
            "Google and GitHub login are available only if their client IDs and secrets are supplied through local configuration or environment variables.",
        ],
    )
    add_heading(document, "2.5 Learning and assessment flow", level=2)
    add_paragraph(
        document,
        "The curriculum is organized into three levels: Beginner, Intermediate, and Advanced. The implementation exposes 15 sections, 30 units, and 90 lessons. A learner who starts from scratch is assigned Beginner immediately. A learner who chooses the placement test answers 18 questions (6 per level bucket), and the resulting score sets the stored skill level. Access to higher-level content is gated by level and by completion of the previous level."
    )
    add_bullets(
        document,
        [
            "Lesson completion grants XP once and stores a `LessonProgress` record.",
            "Lesson quizzes can be passed independently and award XP once.",
            "Unit quizzes are distinct from lesson quizzes and contribute to section readiness.",
            "Section final exams are the certification gate and create certificate records when passed.",
            "Projects are no longer available as early as mere section visibility; they require the section’s lessons to be completed and the unit quizzes to be passed before submission becomes eligible.",
        ],
    )
    add_heading(document, "2.6 Gamification and competition", level=2)
    add_paragraph(
        document,
        "Gamification is one of the most mature parts of the current implementation. UXFriends tracks XP, streaks, unlocked badges, per-day goals, active challenges, and weekly leagues. These systems are backend-driven. The dashboard reads live challenge, streak, reward, and recommendation data from the API rather than hardcoded mock values."
    )
    add_table(
        document,
        ["Mechanic", "Current behavior"],
        [
            ["XP", "Awarded for lessons, quizzes, projects, revision sessions, streak rewards, challenge rewards, and other milestone actions"],
            ["Streaks", "Calculated from qualifying activity days, with milestone rewards at specific day counts"],
            ["Badges", "Large multi-category catalog covering learning, quizzes, projects, revision, streaks, challenges, and XP"],
            ["Daily goals", "Per-user daily goal selection from catalog-based templates, one goal from each category"],
            ["Challenges", "Level-based daily/weekly/personal challenges with join and claim logic"],
            ["Leagues", "Weekly cohort ranking inside level-matched leagues, with Bronze through Diamond tiers and weekly reset behavior"],
        ],
    )
    add_heading(document, "2.7 Portfolio and profile", level=2)
    add_bullets(
        document,
        [
            "Portfolio privacy is now only public or private. The previously discussed link-only mode was removed from the live product flow.",
            "Completed approved projects appear in the user portfolio and can be seen by others if the portfolio is public.",
            "The leaderboard allows users to click another user’s name or avatar and open the public portfolio page when privacy allows it.",
            "The profile page shows professional title, current level, joined date, bio, progress, XP, streak, badges, and section/learning progress summaries.",
        ],
    )
    add_heading(document, "2.8 AI coach and revision support", level=2)
    add_paragraph(
        document,
        "The AI layer currently supports coaching, analysis, revision planning, and persisted chat continuity. It does not yet implement a truly autonomous content-authoring pipeline that generates full quizzes or formal monthly report documents. Instead, it works with learner data already stored in the backend such as lesson progress, quiz attempts, XP logs, and challenge history."
    )
    add_bullets(
        document,
        [
            "The AI coach page shows recommendations, streak context, review priorities, and revision plan summaries.",
            "The analysis page uses real mistake history from quiz attempts and deduplicates repeated wrong answers.",
            "Revision plans are refreshable and map review topics onto the next study days based on current mistakes.",
            "Chat sessions are stored per user in the backend and trimmed to a rolling window of the latest 16 messages.",
            "The FastAPI service can operate in heuristic mode or local-model mode. The local model path is runtime-configurable and currently defaults to lightweight local models rather than a heavy large model.",
        ],
    )

    add_heading(document, "Chapter 3: Data, APIs, and Business Rules", level=1)
    add_heading(document, "3.1 Database provider and migration model", level=2)
    add_paragraph(
        document,
        "The current implementation uses EF Core with SQLite as the default local provider. SQL Server remains supported through configuration, but it is not the current default used by the working application. A key recent change is that local SQLite startup was cleaned up to become migration-based rather than relying on `EnsureCreated` as the primary source of schema truth."
    )
    add_heading(document, "3.2 Current data snapshot", level=2)
    add_table(
        document,
        ["Seeded catalog area", "Current quantity"],
        [
            ["Levels", "3"],
            ["Sections", "15"],
            ["Units", "30"],
            ["Lessons", "90"],
            ["Projects", "15"],
            ["Certificates", "15"],
            ["Challenges", "36"],
            ["Badges", "49"],
            ["Total course hours", "73.9 hours"],
        ],
    )
    add_heading(document, "3.3 Core persisted entities", level=2)
    add_bullets(
        document,
        [
            "Identity and account state: `AspNetUsers`, `PasswordResetTokens`, `EmailVerificationTokens`.",
            "Learning structure and progress: `Sections`, `Units`, `Lessons`, `LessonProgress`, `Quizzes`, `Questions`, `QuizAttempts`, `QuizAttemptAnswers`, `UserNotes`.",
            "Project and certificate data: `Projects`, `ProjectSubmissions`, `Certificates`.",
            "Gamification data: `XPLogs`, `Badges`, `UserBadges`, `Challenges`, `UserChallenges`.",
            "AI continuity: `AIChatSessions`.",
        ],
    )
    add_heading(document, "3.4 Important business rules", level=2)
    add_table(
        document,
        ["Business rule", "Current implemented rule"],
        [
            ["Placement test", "One-time only; 18 questions; score determines stored level"],
            ["Level progression", "Higher levels unlock only after previous-level certification requirements are satisfied"],
            ["Project eligibility", "Requires section lesson completion and passed unit quizzes before submission"],
            ["Project review", "Auto-approved immediately; no human reviewer flow currently exists"],
            ["Certificates", "Issued when the learner passes the section final exam"],
            ["Challenge reward", "Cannot be claimed until progress reaches target and current cycle remains valid"],
            ["League cohort", "Weekly level-matched cohort remains stable through the league week and resets with weekly results"],
            ["Email verification", "Optional but time-bounded; unverified accounts can expire after 48 hours"],
        ],
    )
    add_heading(document, "3.5 Integrations and communication", level=2)
    add_bullets(
        document,
        [
            "SMTP integration is used for password reset and email verification.",
            "External authentication providers are Google and GitHub, each using a backend callback and short-lived code exchange flow.",
            "The AI service communicates over HTTP with the frontend through a dedicated coach API client and can call local-model infrastructure through an Ollama-compatible endpoint.",
            "The public landing page reads live platform totals from the backend, which means the displayed course hours and lesson counts can evolve with the real seed catalog.",
        ],
    )

    add_heading(document, "Chapter 4: Feature Status Matrix", level=1)
    add_heading(document, "4.1 Implemented features", level=2)
    add_table(
        document,
        ["Feature", "Status", "Notes"],
        [
            ["Register / login / logout", "Implemented", "Cookie-based authentication and protected routes are working"],
            ["Forgot password via SMTP", "Implemented", "Email code flow with reset validation"],
            ["Optional email verification", "Implemented", "Verification banner, send flow, token persistence, 48-hour policy"],
            ["Placement onboarding test", "Implemented", "One-time assessment with timer and level assignment"],
            ["Structured lessons and quizzes", "Implemented", "Sections, units, lessons, lesson/unit/section assessments"],
            ["Lesson notes", "Implemented", "Saved in database and reopened later"],
            ["Projects and portfolio publishing", "Implemented", "Project submission auto-approves and appears in portfolio"],
            ["Daily goals / challenges / leagues", "Implemented", "Backend-driven and visible in dashboard/challenges pages"],
            ["AI coach / analysis / revision", "Implemented", "Current AI features are based on progress and mistakes"],
        ],
    )
    add_heading(document, "4.2 Removed or de-scoped features", level=2)
    add_table(
        document,
        ["Old scope item", "Current status", "Reason / impact"],
        [
            ["Shop, transactions, and coins economy", "Removed", "No longer part of the active product flow"],
            ["Link-only portfolio privacy", "Removed from product flow", "Portfolio is now public or private only"],
            ["Human review workflow for projects", "Not implemented", "Projects are auto-approved in the current release"],
            ["Admin/teacher content management portal", "Not implemented", "No admin UI exists in the shipped application"],
        ],
    )
    add_heading(document, "4.3 Planned or not yet fully implemented", level=2)
    add_bullets(
        document,
        [
            "A richer real-model AI experience exists as a runtime option, but the coaching quality still depends on local model availability and configuration.",
            "A dedicated production deployment topology is not yet the source of truth for the local working app, even though the architecture can later be hosted on cloud infrastructure.",
            "Formal reviewer-based project moderation, admin content authoring, and teacher dashboards remain future work rather than shipped features.",
            "Fully autonomous AI-generated quizzes and formal weekly/monthly report artifacts are not the present implemented behavior.",
        ],
    )
    add_heading(document, "4.4 Frontend/backend alignment notes", level=2)
    add_bullets(
        document,
        [
            "The frontend now relies on backend truth for landing metrics, profile state, streaks, rewards, notes, and challenge status rather than hardcoded mock data.",
            "Some backend remnants still exist for deprecated portfolio access-key behavior, but the active frontend no longer exposes link-only privacy.",
            "The AI coach frontend supports persisted conversations and works against the backend-stored session model rather than only browser memory.",
        ],
    )

    add_heading(document, "Chapter 5: Quality, Operations, and Known Limitations", level=1)
    add_heading(document, "5.1 Current quality and test posture", level=2)
    add_paragraph(
        document,
        "The application has been repeatedly validated through live manual testing, targeted API checks, database resets, and build verification. The main runtime services are the React frontend on port 5173, the ASP.NET Core backend on port 5176, and the FastAPI AI service on port 8001. Current documentation and QA records were updated after resetting the local database and verifying that the fresh system still seeds its learning catalog correctly."
    )
    add_bullets(
        document,
        [
            "Frontend build passes successfully under Vite.",
            "Backend build passes successfully under .NET 8.",
            "The AI FastAPI service can be launched independently and exposes a `/health` endpoint.",
            "The local database can be reset safely and recreated from the current migration-aligned startup flow.",
        ],
    )
    add_heading(document, "5.2 Current security hardening state", level=2)
    add_bullets(
        document,
        [
            "JWT tokens are no longer stored in browser local storage; the application uses an `HttpOnly` auth cookie.",
            "Tracked shared configuration contains placeholders, while real secrets should be kept in local-only development configuration or environment variables.",
            "Email verification and password reset both use persisted token records rather than trusting client state alone.",
            "OAuth callbacks are configured through backend public base URL settings instead of loosely depending on incoming host values.",
        ],
    )
    add_heading(document, "5.3 Known limitations", level=2)
    add_bullets(
        document,
        [
            "The AI quality still varies depending on whether the service runs in heuristic mode or local-model mode.",
            "Project approval is automatic because no reviewer/admin moderation workflow exists yet.",
            "The current product is optimized for the implemented web platform; a separate native mobile application does not exist.",
            "Some secondary backend remnants from earlier product ideas still exist in code, even when the active frontend flow no longer uses them.",
        ],
    )
    add_heading(document, "5.4 Recommended future improvements", level=2)
    add_bullets(
        document,
        [
            "Introduce a true content-management or admin authoring workflow if curriculum editing needs to move out of seed code.",
            "Add reviewer moderation for projects if portfolio quality assurance becomes a product requirement.",
            "Expand AI coach quality with stronger local or hosted models while preserving the current persisted mistake-bank context.",
            "Introduce automated integration and end-to-end tests to reduce regression risk across auth, learning, AI, and gamification flows.",
        ],
    )

    add_heading(document, "Chapter 6: Diagram and Graduation Guidance", level=1)
    add_heading(document, "6.1 How this document relates to the diagram work", level=2)
    add_paragraph(
        document,
        "The updated project book should be used together with the implementation-derived diagram guide generated from the live codebase. That guide is stored in `docs/system-diagram-guide.md` and includes Mermaid and PlantUML material for class diagrams, ER diagrams, use cases, activity flows, and sequence diagrams."
    )
    add_heading(document, "6.2 Recommended diagram split", level=2)
    add_bullets(
        document,
        [
            "Auth and account lifecycle",
            "Learning content and assessment",
            "Gamification, portfolio, and AI support",
            "Infrastructure and integration boundaries",
        ],
    )
    add_heading(document, "6.3 Graduation defense notes", level=2)
    add_bullets(
        document,
        [
            "Defend the system as a real implemented educational platform, not as a speculative architecture.",
            "Be explicit about what is implemented today versus what remains future work.",
            "Explain why some product decisions changed during implementation, such as removing coins/shop and simplifying portfolio privacy.",
            "Clarify that SQLite is the current active local provider while SQL Server remains an optional future deployment choice, not the current default runtime truth.",
        ],
    )

    document.add_page_break()
    add_heading(document, "Appendix A: Current Backend Modules", level=1)
    add_table(
        document,
        ["Category", "Current implemented modules"],
        [
            ["Controllers", "AI, Auth, Certificates, Challenges, Dashboard, Leaderboard, Lessons, Portfolio, Projects, Public, Quizzes, Search, Sections, Units, Users"],
            ["Support and policy classes", "LearningAccessPolicy, DailyGoalEngine, StreakEngine, LeagueEngine, BadgeSyncEngine, ChallengeEngine, ChallengeProgressEngine, ChallengeRewardSettlementEngine, GamificationHelper, EmailVerificationPolicy, LearningExperienceCatalog"],
            ["Infrastructure", "UXFriendsDbContext, DatabaseSeeder, CatalogContent, CourseBlueprintCatalog, CourseBlueprintExpansion"],
            ["Tokens and mail", "PasswordResetToken, EmailVerificationToken, EmailService"],
        ],
    )

    add_heading(document, "Appendix B: Current Frontend Pages and Route Groups", level=1)
    add_table(
        document,
        ["Area", "Representative pages"],
        [
            ["Public and auth", "Landing, Login, Signup, Forgot Password, Reset Password, Verify Email, OAuth callback"],
            ["Onboarding", "Onboarding Welcome, Placement Assessment"],
            ["Learning", "Learn, Unit Overview, Lesson Player, Quiz, Unit Quiz, Section Quiz, Section Quiz Result"],
            ["Projects and portfolio", "Projects List, Project Brief, Project Submission, Project Status, Portfolio, Public Portfolio, Privacy Settings"],
            ["AI and revision", "AI Coach Hub, AI Chat, AI Analysis, Recommendations, Revision Center, Revision Plan, Revision Session, Revision Result"],
            ["Gamification and profile", "Dashboard, Challenges Hub, Challenge Detail, Leaderboard, Profile, Settings, Share Page"],
        ],
    )

    add_heading(document, "Appendix C: Added, Removed, and Deferred Scope", level=1)
    add_heading(document, "Added during implementation", level=2)
    add_bullets(
        document,
        [
            "Cookie-based auth with safer frontend storage behavior",
            "Optional email verification with timed expiry rule",
            "Persisted AI chat sessions",
            "Lesson notes saved to the database",
            "Public portfolio viewing from leaderboards",
            "Weekly leagues and daily goals with backend-driven logic",
        ],
    )
    add_heading(document, "Removed or simplified during implementation", level=2)
    add_bullets(
        document,
        [
            "Shop/transaction flow",
            "Coins as an active user-facing reward mechanic",
            "Link-only portfolio privacy",
            "Reviewer moderation requirement for projects",
        ],
    )
    add_heading(document, "Deferred future scope", level=2)
    add_bullets(
        document,
        [
            "Admin content-management console",
            "Teacher/reviewer workflow",
            "Fully autonomous AI quiz authoring and formal generated monthly report documents",
            "Production cloud deployment blueprint as the actively running environment",
        ],
    )


def add_footer(document: Document) -> None:
    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("UXFriends • Updated Project Documentation • April 2026")
        run.font.size = Pt(9)
        run.font.color.rgb = GRAY


def build_document() -> Path:
    document = setup_document()
    add_cover_page(document)
    add_manual_toc(document)
    add_content(document)
    add_footer(document)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    output = build_document()
    print(output)
