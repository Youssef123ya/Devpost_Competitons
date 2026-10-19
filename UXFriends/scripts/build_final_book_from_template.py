from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TEMPLATE_PATH = Path("/Users/hendawiy/Downloads/untitled folder/GP-Template (update).docx")
OUTPUT_PATH = Path("/Users/hendawiy/Downloads/untitled folder/Final final Book.docx")

NUB_BLUE = RGBColor(0x06, 0x68, 0xC4)


def strip_document_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_run_font(run, size: float | None = None, bold: bool | None = None, italic: bool | None = None, color=None):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_paragraph(
    doc: Document,
    text: str = "",
    *,
    style: str = "Normal",
    align: int | None = None,
    before: float | None = None,
    after: float | None = None,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color=None,
):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_bullets(doc: Document, items: list[str], *, style: str = "Normal", indent: float = 0.25):
    for item in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.left_indent = Inches(indent)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(f"- {item}")
        set_run_font(run, size=12)


def add_simple_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_run_font(r, size=11, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            r = p.add_run(value)
            set_run_font(r, size=11)
    doc.add_paragraph("")


def extract_template_images() -> tuple[bytes, bytes]:
    with ZipFile(TEMPLATE_PATH) as zf:
        left = zf.read("word/media/image1.png")
        right = zf.read("word/media/image2.png")
    return left, right


def add_cover(doc: Document) -> None:
    left_logo, right_logo = extract_template_images()

    top = doc.add_table(rows=1, cols=3)
    top.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in top.rows:
        row.cells[0].width = Inches(1.7)
        row.cells[1].width = Inches(3.9)
        row.cells[2].width = Inches(1.7)
    for cell in top.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    left_p = top.cell(0, 0).paragraphs[0]
    left_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    left_p.add_run().add_picture(BytesIO(left_logo), width=Inches(1.35))

    center = top.cell(0, 1)
    for line in [
        "Nahda University in Beni Suef",
        "Faculty of Computer Science",
        "Computer science department",
    ]:
        p = center.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        set_run_font(r, size=18, bold=True, color=NUB_BLUE)
    if center.paragraphs and center.paragraphs[0].text == "":
        center._tc.remove(center.paragraphs[0]._p)

    right_p = top.cell(0, 2).paragraphs[0]
    right_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    right_p.add_run().add_picture(BytesIO(right_logo), width=Inches(1.45))

    add_paragraph(doc, "", before=10)
    add_paragraph(doc, "UXFriends", align=WD_ALIGN_PARAGRAPH.CENTER, size=22, bold=True, before=22, after=18)

    for student in [
        "Mohaned Mohamed Nabil",
        "Youssef Abdelfattah Mahmoud",
        "Ahmed Abdelrahman Ahmed",
        "Abdelrahman Mohamed Ahmed",
        "Abdelrahman Elgamry Hassan",
    ]:
        add_paragraph(doc, student, align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, after=0)

    add_paragraph(doc, "", after=4)
    add_paragraph(doc, "Supervised By", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, italic=True, after=0)
    add_paragraph(doc, "Dr/ Abouelela Abdou Abouelela", align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, after=0)
    add_paragraph(doc, "Eng/ Shahd Mohamed", align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, after=18)

    add_paragraph(
        doc,
        "Submitted for the partial fulfillment of Bachelor of Computer Science degree.",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=14,
        italic=True,
        before=26,
        after=14,
    )
    add_paragraph(doc, "APR, 2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=18, bold=True, after=18)
    add_paragraph(doc, "ABSTRACT", align=WD_ALIGN_PARAGRAPH.CENTER, size=18, bold=True, after=10)

    abstract_paragraphs = [
        "UXFriends is a web-based UX/UI learning platform that transforms a fragmented self-learning experience into a structured pathway that moves from beginner foundations to advanced topics. The system organizes the curriculum into levels, sections, units, and lessons, then reinforces learning through quizzes, section projects, certificates, and portfolio publishing.",
        "The implemented platform combines a React and Vite frontend, an ASP.NET Core backend API, EF Core persistence with SQLite as the current default local database provider, and a FastAPI-based AI coach service. The application supports cookie-based authentication, optional email verification, SMTP password reset, onboarding level determination, lesson notes, mistake-driven revision planning, and public/private portfolio sharing.",
        "To improve motivation and retention, UXFriends includes XP, streaks, badges, daily goals, active challenges, and weekly leagues. The current implementation was validated through live workflow testing, backend API verification, database reset and reseeding, and integration checks across authentication, learning, portfolio, gamification, and AI support flows.",
    ]
    for text in abstract_paragraphs:
        p = add_paragraph(doc, "", style="Normal", after=8)
        p.paragraph_format.first_line_indent = Inches(0.35)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(text)
        set_run_font(r, size=12)

    doc.add_page_break()


def add_declaration(doc: Document) -> None:
    add_paragraph(doc, "DECLARATION", style="Heading 1", align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=12)
    add_paragraph(
        doc,
        "We hereby declare that this dissertation is entirely our own work and original. We understand that if any plagiarism is discovered at any stage, the group may be assigned a fail grade and the submitted work may be rejected according to university regulations.",
        style="No Spacing",
        after=10,
    )
    add_paragraph(doc, "Group members:", style="No Spacing", bold=True, after=6)
    members = [
        "Mohaned Mohamed Nabil",
        "Youssef Abdelfattah Mahmoud",
        "Ahmed Abdelrahman Ahmed",
        "Abdelrahman Mohamed Ahmed",
        "Abdelrahman Elgamry Hassan",
    ]
    for member in members:
        add_paragraph(doc, f"{member}\t\t___________________", style="No Spacing", after=5)
    add_paragraph(doc, "", after=14)
    sign_table = doc.add_table(rows=1, cols=3)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for text in ["Supervisor", "Department Head", "Vice Dean"]:
        cell = sign_table.rows[0].cells[len(sign_table.rows[0].cells) - 3 + ["Supervisor", "Department Head", "Vice Dean"].index(text)]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(text), size=12, bold=True)
        p.add_run("\n\n…………………………….")
    doc.add_page_break()


def add_acknowledgment(doc: Document) -> None:
    add_paragraph(doc, "ACKNOWLEDGMENT", style="Heading 1", align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=12)
    text = (
        "We would like to express our sincere gratitude to our supervisors, Dr. Abouelela Abdou Abouelela and "
        "Eng. Shahd Mohamed, for their guidance, feedback, and continuous support throughout the development of UXFriends. "
        "We also thank the Faculty of Computer Science at Nahda University in Beni Suef for providing the academic environment "
        "that made this work possible. Finally, we appreciate our teammates and families for their encouragement during the analysis, design, implementation, and testing stages of the project."
    )
    p = add_paragraph(doc, "", after=8)
    p.paragraph_format.first_line_indent = Inches(0.35)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_run_font(p.add_run(text), size=12)
    doc.add_page_break()


def add_contents(doc: Document) -> None:
    add_paragraph(doc, "TABLE OF CONTENTS", style="No Spacing", align=WD_ALIGN_PARAGRAPH.CENTER, size=18, bold=True, after=10)
    toc_lines = [
        "Chapter 1: Introduction",
        "Chapter 2: Background / Existing Work",
        "Chapter 3: Software Requirements Specification",
        "Chapter 4: Software Design",
        "Chapter 5: System Implementation & Validation",
        "Chapter 6: Conclusion & Future Work",
        "Appendices",
        "References",
    ]
    for line in toc_lines:
        add_paragraph(doc, line, style="Normal", after=4)
    add_paragraph(doc, "", after=8)
    add_paragraph(doc, "LIST OF FIGURES", style="No Spacing", align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, after=8)
    figures = [
        "Figure 4.1 Current UXFriends runtime architecture",
        "Figure 4.2 Learning and progression activity flow",
        "Figure 4.3 Use case diagram",
        "Figure 4.4 Class diagram",
        "Figure 4.5 Sequence diagram",
        "Figure 4.6 Entity relationship diagram",
    ]
    for fig in figures:
        add_paragraph(doc, fig, style="List Paragraph", after=2)
    add_paragraph(doc, "", after=6)
    add_paragraph(doc, "LIST OF TABLES", style="No Spacing", align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, after=8)
    tables = [
        "Table 2.1 Comparison of UX/UI learning platforms",
        "Table 3.1 Main product functions",
        "Table 5.1 Validation checklist and observed result",
    ]
    for tbl in tables:
        add_paragraph(doc, tbl, style="List Paragraph", after=2)
    doc.add_page_break()


def chapter_title(doc: Document, chapter_no: int, title: str) -> None:
    add_paragraph(doc, f"CHAPTER {chapter_no}", style="Normal", align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, before=6, after=2)
    add_paragraph(doc, title, style="Normal", align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, after=16)


def section_heading(doc: Document, text: str) -> None:
    add_paragraph(doc, text, style="Heading 1", after=8)


def subsection_heading(doc: Document, text: str) -> None:
    add_paragraph(doc, text, style="Heading 2", after=6)


def body_paragraph(doc: Document, text: str) -> None:
    p = add_paragraph(doc, "", style="Normal", after=8)
    p.paragraph_format.first_line_indent = Inches(0.35)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_run_font(p.add_run(text), size=12)


def add_chapter_1(doc: Document) -> None:
    chapter_title(doc, 1, "INTRODUCTION")
    body_paragraph(doc, "Learning UX/UI design has become more accessible than ever, yet accessibility has not automatically produced a coherent learning journey. Many beginners move between scattered videos, articles, and short courses without a clear pathway, practical reinforcement, or timely feedback. UXFriends was developed to solve this educational gap by delivering a structured and engaging web platform that combines guided learning, assessment, practice, motivation, and coaching in one place.")

    subsection_heading(doc, "1.1 Problem domain")
    body_paragraph(doc, "The problem domain of UXFriends is educational technology with a focus on UX/UI design training. The system is designed for learners who need a curriculum that starts with foundations and gradually builds toward stronger design analysis, visual communication, interaction design, accessibility awareness, and project readiness. This domain requires both conceptual learning and repeated practical reinforcement.")

    subsection_heading(doc, "1.2 Problem statement")
    body_paragraph(doc, "Beginners in UX/UI often suffer from four connected problems: fragmented learning materials, weak practice discipline, low long-term motivation, and unclear awareness of strengths and weaknesses. Existing resources may teach theory well, but they do not always provide a single pathway that combines sequential content, projects, measurable progress, and revision support. As a result, many learners stop early or fail to translate theory into practical design decisions.")

    subsection_heading(doc, "1.3 Proposed system")
    body_paragraph(doc, "UXFriends is a gamified UX/UI learning platform implemented as a web application. It organizes the learning journey into three levels—Beginner, Intermediate, and Advanced—and each level contains sections, units, and lessons. The system also integrates quizzes, section final exams, capstone projects, certificates, public portfolios, AI coaching, and progress-based gamification.")

    subsection_heading(doc, "1.3.1 Aims and Objectives")
    add_bullets(doc, [
        "Provide a structured learning experience from foundation to advanced topics.",
        "Give learners practical reinforcement through quizzes, notes, and section projects.",
        "Improve consistency with XP, streaks, badges, daily goals, challenges, and weekly leagues.",
        "Support self-correction through AI-driven recommendations, mistake analysis, and revision planning.",
        "Help learners build visible outcomes through certificates and portfolio-ready projects.",
    ])

    subsection_heading(doc, "1.3.2 Proposed system features")
    add_bullets(doc, [
        "Cookie-based registration, login, logout, forgot password, and optional email verification.",
        "One-time onboarding placement test or direct beginner start-from-scratch flow.",
        "Level-based learning access across 15 sections, 30 units, and 90 lessons.",
        "Lesson quizzes, unit quizzes, section final exams, and section certificates.",
        "Lesson notes stored in the database for later review.",
        "Auto-approved capstone project submissions that publish into the user portfolio.",
        "AI coach, AI analysis, revision plan, revision sessions, and persisted chat history.",
        "Backend-driven leaderboard, leagues, challenges, and daily goals.",
    ])

    subsection_heading(doc, "1.4 Project Methodology")
    body_paragraph(doc, "The implementation of UXFriends followed an iterative development approach inspired by Agile and Scrum principles. Features were delivered in layers, starting from authentication and learning navigation, then expanding into progress tracking, projects, certificates, AI support, and security hardening. This approach allowed the product scope to evolve based on implementation realities while preserving the educational goals of the platform.")

    subsection_heading(doc, "1.4.1 Development phases")
    add_bullets(doc, [
        "Requirements analysis and scope definition.",
        "UI/UX planning and route/component decomposition.",
        "Backend API and EF Core data model implementation.",
        "Curriculum, quizzes, projects, and gamification logic implementation.",
        "AI coaching service integration and persistence support.",
        "Testing, bug fixing, security hardening, and documentation updates.",
    ])

    subsection_heading(doc, "1.4.2 Key methodology principles")
    add_bullets(doc, [
        "Iterative delivery with real end-to-end verification.",
        "Backend as the source of truth for user progress and reward state.",
        "Separation of concerns between frontend, backend, and AI service.",
        "Continuous refinement of product scope according to implemented behavior.",
    ])

    subsection_heading(doc, "1.5 Resource Requirement")
    subsection_heading(doc, "1.5.1 Human Resources")
    body_paragraph(doc, "The project was produced by a five-student team under academic supervision. The team members are Mohaned Mohamed Nabil, Youssef Abdelfattah Mahmoud, Ahmed Abdelrahman Ahmed, Abdelrahman Mohamed Ahmed, and Abdelrahman Elgamry Hassan. Academic guidance was provided by Dr. Abouelela Abdou Abouelela and Eng. Shahd Mohamed.")

    subsection_heading(doc, "1.5.2 Software Resources")
    add_bullets(doc, [
        "React 19, Vite, and TypeScript for the frontend.",
        "ASP.NET Core 8 and ASP.NET Identity for the backend.",
        "EF Core with SQLite as the default local provider and SQL Server as an optional provider.",
        "FastAPI for the AI coach service.",
        "SMTP integration for email verification and password reset.",
        "Optional Ollama-compatible local model runtime for AI local mode.",
    ])

    subsection_heading(doc, "1.5.3 Hardware and Infrastructure Resources")
    body_paragraph(doc, "The current system is designed as a browser-based application that can run on standard desktops, laptops, tablets, and modern smartphones. Development and local testing require Node.js, the .NET SDK, Python, and optional AI model tooling. The runtime architecture is service-based, which also prepares the system for future hosted deployment if needed.")

    subsection_heading(doc, "1.5.4 Economic Resources")
    body_paragraph(doc, "The project relies primarily on student development effort and open-source tooling. The local development environment uses freely available frameworks and libraries. Optional external costs may later arise from SMTP hosting, domain hosting, database hosting, or hosted AI APIs, but the current implementation is able to function in a local environment with minimal cost.")

    subsection_heading(doc, "1.6 Report Layout")
    body_paragraph(doc, "This report begins with the project introduction and problem analysis, then reviews existing work, then describes the software requirements and design decisions, and finally documents the implementation, validation results, conclusion, and future work.")
    doc.add_page_break()


def add_chapter_2(doc: Document) -> None:
    chapter_title(doc, 2, "BACKGROUND / EXISTING WORK")
    subsection_heading(doc, "2.0 Introduction")
    body_paragraph(doc, "This chapter positions UXFriends within the context of modern online design learning. It reviews how other platforms approach UX/UI education, where they succeed, and where they leave gaps that UXFriends attempts to address through integrated progression, gamification, and practical output.")

    subsection_heading(doc, "2.1 Overview of existing projects")
    body_paragraph(doc, "The UX/UI learning market includes broad course marketplaces, specialist design education platforms, and gamified microlearning tools. Coursera and Udemy provide broad catalogs with many instructors and topics. Interaction Design Foundation offers structured theory-rich UX education. UXcel provides short, gamified UX lessons and quick assessments. These platforms show that there is strong demand for guided design learning, but they differ in depth, interactivity, and personalization.")
    add_simple_table(
        doc,
        ["Platform", "Strengths", "Limitations relative to UXFriends"],
        [
            ["Coursera / Udemy", "Large course catalogs and recognizable instructors", "Mostly passive learning and weak integrated gamification"],
            ["Interaction Design Foundation", "Structured UX-focused curriculum", "Less emphasis on continuous competitive motivation and practical in-app gamification"],
            ["UXcel", "Gamified bite-sized UX learning", "Limited control over portfolio/project flow and currently different personalization direction"],
        ],
    )

    subsection_heading(doc, "2.2 Limitations of existing projects")
    body_paragraph(doc, "Many existing solutions separate content from long-term motivation and separate learning theory from project tracking. Beginners may complete scattered lessons without understanding their overall level, without building visible work, and without seeing precise mistake patterns. Some platforms also emphasize passive reading more than repeated correction and revision.")

    subsection_heading(doc, "2.3 Innovations of our project")
    body_paragraph(doc, "UXFriends differentiates itself by combining structured progression, portfolio publishing, mistake-based AI revision support, and backend-driven gamification in one application. The current system includes level placement, section certification, project eligibility rules, league competition, public portfolios, persisted lesson notes, and AI chat continuity. These features make the platform feel less like an isolated course viewer and more like an active learning environment.")
    doc.add_page_break()


def add_chapter_3(doc: Document) -> None:
    chapter_title(doc, 3, "SOFTWARE REQUIREMENTS SPECIFICATION")
    section_heading(doc, "3.1 Introduction")
    subsection_heading(doc, "3.1.1 Purpose")
    body_paragraph(doc, "This Software Requirements Specification defines the implemented scope of UXFriends as a gamified UX/UI learning platform. The document covers the current web application behavior across frontend, backend, database, and AI support components.")

    subsection_heading(doc, "3.1.2 Document Conventions")
    body_paragraph(doc, "The report uses standard chapter and subsection numbering. Functional behaviors are described according to the implemented system rather than speculative future features. When a feature is planned but not yet implemented, it is explicitly marked as future work or deferred scope.")

    subsection_heading(doc, "3.1.3 Intended Audience and Reading Suggestions")
    body_paragraph(doc, "This specification is intended for examiners, software developers, QA reviewers, graduation-project supervisors, and future maintainers of the UXFriends codebase. Readers interested in product scope should focus on Sections 3.2 to 3.4, while technical reviewers should also examine Chapters 4 and 5.")

    subsection_heading(doc, "3.1.4 Product Scope")
    body_paragraph(doc, "UXFriends is a browser-based educational platform for learning UX/UI design. It supports structured content, assessments, projects, gamification, AI-assisted revision, portfolio publishing, and secure account management. The current platform emphasizes the learner role and does not yet include a complete admin or teacher management console.")

    section_heading(doc, "3.2 Overall Description")
    subsection_heading(doc, "3.2.1 Product Perspective")
    body_paragraph(doc, "The application is implemented as a multi-service system: a React/Vite frontend communicates with an ASP.NET Core backend API, while a FastAPI service supports AI coaching features. The backend stores learning, user, assessment, and gamification data through EF Core.")

    subsection_heading(doc, "3.2.2 Product Functions")
    add_simple_table(
        doc,
        ["Function area", "Current implemented behavior"],
        [
            ["Account management", "Register, login, logout, forgot/reset password, optional email verification, Google/GitHub OAuth"],
            ["Learning flow", "Browse levels, sections, units, lessons, complete lessons, store notes, take quizzes and exams"],
            ["Projects and portfolio", "View project briefs, submit projects, auto-approve, publish approved work to portfolio"],
            ["Gamification", "Track XP, streaks, badges, daily goals, challenges, and weekly leagues"],
            ["AI support", "View analysis, recommendations, revision plan, revision sessions, and persisted coaching chat"],
        ],
    )

    subsection_heading(doc, "3.2.3 User Classes and Characteristics")
    add_bullets(doc, [
        "Guest: can view the landing page and public portfolios, and can start registration or login.",
        "Learner/User: can consume content, complete quizzes and projects, view profile and portfolio, and use AI coaching features.",
        "External systems: SMTP mail provider, OAuth providers, and optional local AI model runtime.",
    ])

    subsection_heading(doc, "3.2.4 Operating Environment")
    body_paragraph(doc, "The frontend runs on modern browsers. The backend runs on ASP.NET Core 8. The AI service runs on Python with FastAPI. The local development environment currently uses SQLite by default, while SQL Server remains an optional configured provider.")

    subsection_heading(doc, "3.2.5 Design and Implementation Constraints")
    add_bullets(doc, [
        "The platform currently focuses on learners only and does not depend on a reviewer role for project approval.",
        "The AI service must remain usable even when a local model is not available, so heuristic mode exists as a fallback.",
        "Portfolio privacy is limited to public or private, and link-only behavior is not part of the active user flow.",
        "Challenge and league logic must stay consistent with level-based grouping and weekly cycle rules.",
    ])

    subsection_heading(doc, "3.2.6 User Documentation")
    body_paragraph(doc, "The current codebase includes a root README for setup, a system diagram guide for graduation documentation, and in-app interface text across dashboard, learning, AI, and profile modules.")

    subsection_heading(doc, "3.2.7 Assumptions and Dependencies")
    body_paragraph(doc, "The application assumes valid SMTP credentials when email flows are enabled, valid OAuth provider credentials when social login is enabled, and an available FastAPI service when AI chat is used. The application also assumes that the backend remains the source of truth for progress and reward calculations.")

    section_heading(doc, "3.3 External Interface Requirements")
    subsection_heading(doc, "3.3.1 User Interfaces")
    body_paragraph(doc, "The application provides responsive interfaces for landing, authentication, dashboard, learning content, quizzes, projects, AI coaching, challenges, leaderboard, profile, portfolio, and settings. The frontend is bilingual, with Arabic as the primary default route and English available through route language parameters.")

    subsection_heading(doc, "3.3.2 Hardware Interfaces")
    body_paragraph(doc, "No specialized hardware is required beyond standard computing devices capable of running a web browser. Optional local AI mode requires a device that can run an Ollama-compatible model service.")

    subsection_heading(doc, "3.3.3 Software Interfaces")
    add_bullets(doc, [
        "REST API integration between React frontend and ASP.NET Core backend.",
        "SMTP integration for verification and password reset emails.",
        "Google and GitHub OAuth flows through backend callback endpoints.",
        "FastAPI chat service for AI coaching features.",
    ])

    subsection_heading(doc, "3.3.4 Communications Interfaces")
    body_paragraph(doc, "The system communicates over HTTP in local development. Authentication is handled through `HttpOnly` cookies. JSON is used for data exchange between frontend and backend, and between the frontend AI client and the AI service.")

    section_heading(doc, "3.4 System Features")
    subsection_heading(doc, "3.4.1 Learning Content and Progress")
    add_bullets(doc, [
        "The system shall provide 3 learning levels, 15 sections, 30 units, and 90 lessons.",
        "The system shall save lesson completion and prevent duplicate XP on repeated completion.",
        "The system shall allow learners to add and update notes for each lesson.",
        "The system shall gate higher-level content according to stored skill level and certification progress.",
    ])

    subsection_heading(doc, "3.4.2 Assessment and Certification")
    add_bullets(doc, [
        "The system shall provide lesson quizzes, unit quizzes, and section final exams.",
        "The system shall store quiz attempts and question-level wrong answers.",
        "The system shall issue section certificates after a successful section final exam.",
        "The onboarding placement test shall be one-time and shall assign the learner level from the score.",
    ])

    subsection_heading(doc, "3.4.3 Projects and Portfolio")
    add_bullets(doc, [
        "The system shall provide one capstone project per section.",
        "The system shall allow project submission only when the section lesson and unit-quiz prerequisites are satisfied.",
        "The current system shall auto-approve valid project submissions.",
        "Approved projects shall appear in the learner portfolio and public portfolio when privacy allows.",
    ])

    subsection_heading(doc, "3.4.4 Gamification and Competition")
    add_bullets(doc, [
        "The system shall track XP, streaks, badges, daily goals, challenges, and weekly leagues.",
        "The system shall calculate challenge progress and reward eligibility from backend truth.",
        "The system shall rank learners weekly inside level-based league cohorts.",
    ])

    subsection_heading(doc, "3.4.5 AI Learning Support")
    add_bullets(doc, [
        "The system shall show strengths, weaknesses, and areas to review based on learner mistakes.",
        "The system shall generate revision plans and revision sessions from real learner progress.",
        "The system shall persist recent AI chat history in the backend.",
        "The AI service shall support heuristic mode and local-model mode.",
    ])

    subsection_heading(doc, "3.4.6 Settings and Account Control")
    add_bullets(doc, [
        "The user shall be able to update profile information and profile photo.",
        "The user shall be able to change password.",
        "The user shall be able to initiate email verification from settings if not already verified.",
        "The user shall be able to set the preferred display language.",
    ])

    section_heading(doc, "3.5 Other Non-functional Requirements")
    subsection_heading(doc, "3.5.1 Performance Requirements")
    add_bullets(doc, [
        "The application shall keep the first interactive state reasonably responsive in production mode through route splitting and cached landing data.",
        "The platform shall support normal educational workflows without blocking on the AI service for unrelated pages.",
    ])

    subsection_heading(doc, "3.5.2 Safety Requirements")
    body_paragraph(doc, "The platform does not control physical devices, but it must avoid destructive data loss through careful account-state handling, token validation, and challenge or reward duplication protection.")

    subsection_heading(doc, "3.5.3 Security Requirements")
    add_bullets(doc, [
        "Authentication shall use `HttpOnly` cookie storage instead of exposing JWTs in browser local storage.",
        "Password reset and email verification shall use persisted tokens with validation and expiration rules.",
        "Protected endpoints shall reject unauthorized access.",
        "Real environment secrets shall remain outside tracked source control configuration.",
    ])

    subsection_heading(doc, "3.5.4 Software Quality Attributes")
    add_bullets(doc, [
        "Maintainability through controller/service separation.",
        "Extensibility through a separate AI microservice and configurable database provider.",
        "Usability through bilingual UI, structured progression, and visual gamification feedback.",
    ])

    subsection_heading(doc, "3.5.5 Business Rules")
    add_bullets(doc, [
        "Unverified accounts may expire after 48 hours from creation.",
        "Projects are auto-approved in the current version because no reviewer flow exists.",
        "Challenge rewards can be claimed only when the target is reached and the cycle remains valid.",
        "The learner level is updated by onboarding result and later by level certification progress.",
    ])
    doc.add_page_break()


def add_chapter_4(doc: Document) -> None:
    chapter_title(doc, 4, "SOFTWARE DESIGN")
    section_heading(doc, "4.1 Design Overview")
    subsection_heading(doc, "4.1.1 Introduction")
    body_paragraph(doc, "UXFriends follows a modular web-application architecture. The frontend is responsible for user interaction, client routing, and localized rendering. The backend owns business rules, persistence, authentication, progress tracking, gamification, and public data delivery. The AI service operates as a dedicated support component for coaching and revision behavior.")

    subsection_heading(doc, "4.1.2 Environment Overview")
    body_paragraph(doc, "In local development, the frontend runs on port 5173, the ASP.NET Core backend runs on port 5176, and the FastAPI AI service runs on port 8001. SQLite is used by default for local persistence, and the database file is stored inside the backend project folder. This environment supports full end-to-end execution of registration, learning, progress, portfolio, and AI workflows.")

    subsection_heading(doc, "4.1.3 System Architecture")
    body_paragraph(doc, "The architecture can be summarized as a three-service model with a shared persistence layer. The learner interacts with the React frontend. The frontend sends authenticated requests to the backend API using cookies. The backend stores user and progress state in EF Core entities. For AI chat and analysis support, the frontend also uses a coach API client that targets the FastAPI service. The AI service operates with current learner context and recent chat history already stored in the backend ecosystem.")

    subsection_heading(doc, "4.1.4 Constraints and Assumptions")
    add_bullets(doc, [
        "The current system assumes learner-facing behavior only, not admin moderation.",
        "The AI service must not block the rest of the application if unavailable.",
        "Backend APIs remain the authoritative source of truth for progress and reward state.",
        "The current report reflects actual implementation, so non-implemented future features are not modeled as active runtime behavior.",
    ])

    section_heading(doc, "4.2 Interfaces and Data Stores")
    subsection_heading(doc, "4.2.1 System Interfaces")
    body_paragraph(doc, "The main interface groups are the learner web interface, backend REST endpoints, SMTP mail delivery, OAuth provider integration, and FastAPI AI chat endpoints. The frontend API client is centralized and sends requests with credentials for auth-protected flows.")

    subsection_heading(doc, "4.2.2 Data Stores")
    body_paragraph(doc, "The main data store is the relational application database managed through EF Core. Persisted data includes users, progress, quiz attempts, wrong answers, notes, projects, submissions, certificates, XP logs, badges, challenge participation, password reset tokens, email verification tokens, and AI chat sessions. The learning catalog is seeded and enriched from code-backed catalog content rather than being authored through an admin CMS.")

    section_heading(doc, "4.3 Structural Design")
    subsection_heading(doc, "4.3.1 Class Diagram")
    body_paragraph(doc, "The core class model includes the user identity entity, curriculum entities such as Section, Unit, Lesson, assessment entities such as Quiz and Question, progress entities such as LessonProgress and QuizAttempt, gamification entities such as XPLog, Badge, UserBadge, Challenge, and UserChallenge, and support entities such as PasswordResetToken, EmailVerificationToken, and AIChatSession. These classes are connected primarily through one-to-many relationships anchored around the learner account.")

    subsection_heading(doc, "4.3.2 Class Descriptions")
    add_bullets(doc, [
        "User is the central identity entity and owns progress, notes, projects, badges, XP logs, and tokens.",
        "Section, Unit, and Lesson model the learning hierarchy.",
        "QuizAttempt and QuizAttemptAnswer capture grading evidence and mistake history.",
        "ProjectSubmission and Certificate represent learner outputs and earned milestones.",
        "AIChatSession stores recent coach conversation state for continuity.",
    ])

    subsection_heading(doc, "4.3.3 Sequence Diagram")
    body_paragraph(doc, "The most important sequences in UXFriends are registration, login, forgot password, lesson completion, quiz submission, project submission, challenge claim, leaderboard retrieval, and AI coach chat. Each of these sequences involves clear frontend-to-backend calls and, in some cases, database writes and optional SMTP or AI provider interactions.")

    subsection_heading(doc, "4.3.4 ER Diagram and Relation Schema")
    body_paragraph(doc, "The entity relationship design is strongly centered around the AspNetUsers table. Most learner-owned records use a user foreign key and follow one-to-many cardinality. The schema also includes many-to-many-like patterns through join-style records such as UserBadges and UserChallenges. The detailed ERD should reflect the currently implemented EF Core model and not older assumptions about removed modules such as shop transactions.")
    doc.add_page_break()


def add_chapter_5(doc: Document) -> None:
    chapter_title(doc, 5, "SYSTEM IMPLEMENTATION & VALIDATION")
    body_paragraph(doc, "This chapter documents how UXFriends was implemented and how the working application was validated across the main user flows. The focus is on real runtime behavior rather than intended architecture only.")

    subsection_heading(doc, "5.1 Frontend Implementation")
    body_paragraph(doc, "The frontend was implemented in React with Vite and TypeScript. The application uses a route-aware shell, central API clients, loading states, language-aware navigation, and page groups for landing, auth, learning, projects, AI, profile, challenges, portfolio, and settings. Data displayed on major pages such as dashboard, profile, challenges, and learning is now fetched from the backend rather than hardcoded.")

    subsection_heading(doc, "5.2 Backend Implementation")
    body_paragraph(doc, "The backend was implemented in ASP.NET Core 8 with controller-based APIs. Controllers were created for authentication, users, sections, units, lessons, quizzes, projects, certificates, dashboard, challenges, leaderboard, portfolio, search, public landing data, and AI support. Business-rule logic is separated into support classes such as LearningAccessPolicy, DailyGoalEngine, StreakEngine, LeagueEngine, BadgeSyncEngine, and ChallengeProgressEngine.")

    subsection_heading(doc, "5.3 AI Service Implementation")
    body_paragraph(doc, "The AI service was implemented with FastAPI. Its role is to support AI coaching, rather than replace the backend as the owner of learner data. The service can work in heuristic mode or local-model mode and is designed to respond quickly to simple interactions while still using backend-derived learner context and rolling message history.")

    subsection_heading(doc, "5.4 Database and Seeding")
    body_paragraph(doc, "The local environment currently uses SQLite as the default active database. The latest startup flow supports clean database resets and migration-aware recreation. Seeded catalog data currently provides 3 levels, 15 sections, 30 units, 90 lessons, 15 projects, 36 challenges, and 49 badges. This makes the development environment immediately usable after a reset.")

    subsection_heading(doc, "5.5 Authentication and Email Flows")
    add_bullets(doc, [
        "Register and login set an HttpOnly authentication cookie.",
        "Forgot password sends an SMTP reset code and validates it before password change.",
        "Email verification stores a token in the database and verifies the account through a clickable link.",
        "Google and GitHub login are supported when provider credentials are configured.",
    ])

    subsection_heading(doc, "5.6 Learning, Projects, and Portfolio Validation")
    body_paragraph(doc, "The learner flow was validated by creating fresh users, progressing through onboarding, completing lessons, storing notes, passing quizzes, unlocking section readiness, submitting projects, and checking portfolio publication. Project submission is intentionally auto-approved in the current implementation because no human moderation workflow exists. Public portfolios were also validated through leaderboard entry navigation.")

    subsection_heading(doc, "5.7 Gamification and Leaderboard Validation")
    body_paragraph(doc, "XP rewards, streak milestones, badges, daily goals, active challenges, and weekly leagues were validated through real user progression and backend checks. Challenge reward claiming depends on backend progress evaluation and valid cycle state. Weekly league ranking is level-based and resets on the weekly cycle, while keeping result information available after the cycle closes.")

    subsection_heading(doc, "5.8 AI and Revision Validation")
    body_paragraph(doc, "The AI-related flows were validated by checking the health endpoint, chat continuity, mistake-based analysis, revision planning, and revision session completion. The AI analysis page uses real wrong-answer history rather than generic placeholders. The chat persistence layer retains recent messages so a learner does not lose all context after a refresh.")

    subsection_heading(doc, "5.9 Validation Checklist and Observed Result")
    add_simple_table(
        doc,
        ["Flow", "Observed result"],
        [
            ["Register / login / logout", "Passed with cookie-based auth"],
            ["Forgot password", "Passed with SMTP-backed reset flow"],
            ["Email verification", "Passed with stored token and verification endpoint"],
            ["Placement onboarding", "Passed with one-time timed test and level storage"],
            ["Lesson completion + notes", "Passed with persisted progress and notes"],
            ["Quiz submission", "Passed with stored attempts and score result"],
            ["Project submission", "Passed with eligibility gating and auto-approval"],
            ["Portfolio privacy", "Passed for public/private behavior"],
            ["Challenges and leagues", "Passed with backend-driven progress and ranking"],
            ["AI coach health and persistence", "Passed with service health and recent-session continuity"],
        ],
    )

    subsection_heading(doc, "5.10 Current Limitations")
    add_bullets(doc, [
        "There is no admin or reviewer dashboard in the current version.",
        "Projects are auto-approved rather than manually reviewed.",
        "AI quality depends on local runtime mode and configuration.",
        "Some deprecated backend remnants still exist even though the main frontend flow no longer uses them.",
    ])
    doc.add_page_break()


def add_chapter_6(doc: Document) -> None:
    chapter_title(doc, 6, "CONCLUSION & FUTURE WORK")
    subsection_heading(doc, "6.1 Conclusion")
    body_paragraph(doc, "UXFriends successfully evolved from an educational concept into a working multi-service learning platform for UX/UI design. The final implemented system now supports structured progression, assessments, projects, portfolio visibility, progress tracking, gamification, email-backed account flows, and AI-assisted revision support. The application demonstrates that a beginner-friendly design-learning experience can be built as an integrated product rather than a collection of disconnected resources.")

    subsection_heading(doc, "6.2 Main Achievements")
    add_bullets(doc, [
        "A complete learner-facing web product was implemented across frontend, backend, AI, and database layers.",
        "The system now includes strong learning-state persistence and real backend-driven gamification.",
        "The product supports practical outputs through certificates and portfolio-ready projects.",
        "Security and operational quality improved through cookie-based auth, email verification, secret cleanup, and migration-aligned database startup.",
    ])

    subsection_heading(doc, "6.3 Current Limitations")
    add_bullets(doc, [
        "No human reviewer or teacher workflow exists yet.",
        "No dedicated admin content-authoring interface currently exists.",
        "Some AI behavior still depends on lightweight or heuristic runtime modes rather than a stronger hosted model.",
        "The production hosting topology is not yet the primary working environment of the current project submission.",
    ])

    subsection_heading(doc, "6.4 Future Work")
    add_bullets(doc, [
        "Build an admin/content-management workflow for curriculum maintenance.",
        "Add reviewer moderation for projects and richer project feedback loops.",
        "Expand the AI coach with stronger model-backed reasoning while preserving current progress context and persistence.",
        "Introduce automated integration and end-to-end testing across auth, learning, portfolio, gamification, and AI flows.",
        "Prepare production-grade deployment and monitoring for hosted public use.",
    ])
    doc.add_page_break()


def add_appendices_and_refs(doc: Document) -> None:
    add_paragraph(doc, "APPENDICES", style="Normal", align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, after=16)
    body_paragraph(doc, "Appendix I may contain the current backend controller inventory, Appendix II may contain the current frontend route groups, and Appendix III may contain the implementation-aligned class, ER, use case, activity, and sequence diagrams prepared from the live codebase.")
    add_paragraph(doc, "REFERENCES", style="Normal", align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, before=18, after=10)
    refs = [
        "[1] React Documentation, \"React,\" https://react.dev/.",
        "[2] Microsoft, \"ASP.NET Core documentation,\" https://learn.microsoft.com/aspnet/core/.",
        "[3] FastAPI Documentation, \"FastAPI,\" https://fastapi.tiangolo.com/.",
        "[4] Microsoft, \"Entity Framework Core,\" https://learn.microsoft.com/ef/core/.",
        "[5] SQLite Documentation, \"SQLite Home Page,\" https://www.sqlite.org/.",
        "[6] UXcel, \"UX learning platform,\" https://uxcel.com/.",
        "[7] Interaction Design Foundation, \"UX and design education,\" https://www.interaction-design.org/.",
        "[8] Ollama Documentation, \"Ollama,\" https://ollama.com/.",
    ]
    for ref in refs:
        add_paragraph(doc, ref, style="No Spacing", after=4)


def set_margins(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)


def build() -> Path:
    doc = Document(TEMPLATE_PATH)
    strip_document_body(doc)
    set_margins(doc)
    add_cover(doc)
    add_declaration(doc)
    add_acknowledgment(doc)
    add_contents(doc)
    add_chapter_1(doc)
    add_chapter_2(doc)
    add_chapter_3(doc)
    add_chapter_4(doc)
    add_chapter_5(doc)
    add_chapter_6(doc)
    add_appendices_and_refs(doc)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build())
